# api_doc_parser.py
import os
import re
import json
from langchain_deepseek import ChatDeepSeek
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from config import DEEPSEEK_API_KEY

# 初始化大模型
llm = ChatDeepSeek(
    model="deepseek-v4-flash",
    api_key=DEEPSEEK_API_KEY,
    temperature=0.1
)


def read_file_content(file_path: str) -> str:
    """读取各种格式的文档内容"""
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt" or ext == ".md":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(file_path)
            lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception as e:
            print(f"读取docx失败: {e}")
            return ""

    elif ext == ".pdf":
        try:
            from PyPDF2 import PdfReader
            text = ""
            for page in PdfReader(file_path).pages:
                if page.extract_text():
                    text += page.extract_text() + "\n"
            return text
        except Exception as e:
            print(f"读取pdf失败: {e}")
            return ""

    else:
        # 其他格式当文本读
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except:
            return ""


def parse_doc_for_selenium(file_path: str) -> dict:
    """
    智能解析文档，提取Selenium测试所需的关键信息
    返回：{url, elements, function_desc}
    """
    # 1. 读取文档内容
    text = read_file_content(file_path)
    if not text:
        return {"url": "", "elements": [], "function_desc": "文档读取失败"}

    # 2. 限制长度，避免超出上下文
    if len(text) > 8000:
        text = text[:8000] + "\n...(截断)"

    # 3. 直接构造prompt字符串（用单花括号，不需要转义）
    system_prompt = """你是一个Web自动化测试专家。请分析以下文档，提取出Selenium自动化测试所需的关键信息。
严格按照JSON格式返回，只返回JSON，不要markdown代码块，不要任何额外解释。

返回格式：
{
  "url": "被测系统的页面地址",
  "elements": [
    {"name": "元素的中文描述", "locator_type": "id或name", "locator_value": "定位值"}
  ],
  "function_desc": "功能描述"
}

注意：
1. URL如果文档里没有，就根据上下文推断一个。
2. 每个输入框、按钮都要提取其id或name属性。
3. 只返回一个有效的JSON对象，不要任何其他文字。"""

    user_prompt = f"文档内容：\n{text}"

    # 4. 调用大模型
    try:
        from langchain_core.messages import SystemMessage, HumanMessage
        response = llm.invoke([
            SystemMessage(content=system_prompt),
            HumanMessage(content=user_prompt)
        ])
        result_text = response.content.strip()

        # 打印原始返回，方便调试
        print("【大模型原始返回】")
        print(result_text[:500])
        print("=" * 50)

        # 5. 三层JSON提取容错
        try:
            result = json.loads(result_text)
        except:
            # 去掉markdown代码块
            cleaned = result_text
            if cleaned.startswith("```json"):
                cleaned = cleaned[7:]
            elif cleaned.startswith("```"):
                cleaned = cleaned[3:]
            if cleaned.endswith("```"):
                cleaned = cleaned[:-3]
            cleaned = cleaned.strip()

            try:
                result = json.loads(cleaned)
            except:
                # 正则提取 { ... }
                import re
                match = re.search(r'\{.*\}', cleaned, re.DOTALL)
                if match:
                    try:
                        result = json.loads(match.group())
                    except:
                        raise ValueError("JSON解析失败")
                else:
                    raise ValueError("未找到JSON内容")

        result.setdefault("url", "")
        result.setdefault("elements", [])
        result.setdefault("function_desc", "未识别功能")
        return result

    except Exception as e:
        print(f"大模型解析失败: {e}")
        return {
            "url": "",
            "elements": [],
            "function_desc": f"解析失败: {str(e)[:50]}"
        }

# 控制台测试
if __name__ == "__main__":
    result = parse_doc_for_selenium("files/登录页面设计文档.txt")
    print(json.dumps(result, ensure_ascii=False, indent=2))
