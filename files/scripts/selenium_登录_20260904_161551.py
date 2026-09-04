#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
登录功能 Selenium 自动化测试脚本
运行方式: pytest test_login.py
"""

import logging
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import (
    NoSuchElementException,
    TimeoutException,
    WebDriverException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.edge.options import Options
from selenium.webdriver.edge.service import Service
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.microsoft import EdgeChromiumDriverManager

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger("login_test")

BASE_URL = "http://127.0.0.1:8080/login"
TESTCASE_FILE = Path(__file__).resolve().parent / "files" / "login_testcases.xlsx"

REQUIRED_FIELDS = [
    "用例ID",
    "账号",
    "密码",
    "场景描述",
    "用例类型",
    "预期结果",
]


def _to_display_text(value):
    """把 Excel 单元格转换为适合界面的文本。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_testcases(excel_path):
    """读取 Excel 测试用例。"""
    if not excel_path.exists():
        raise FileNotFoundError(f"测试用例文件不存在: {excel_path}")

    workbook = load_workbook(excel_path, data_only=True)

    known_sheets = ["登录用例", "登录测试", "用例", "测试用例", "Sheet1", "Sheet"]
    sheet_name = None
    for name in workbook.sheetnames:
        if name in known_sheets:
            sheet_name = name
            break
    if sheet_name is None:
        sheet_name = workbook.sheetnames[0]

    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Excel 文件 {excel_path.name} 中没有数据")

    headers = [_to_display_text(cell) for cell in rows[0]]
    if not all(field in headers for field in REQUIRED_FIELDS):
        raise ValueError(
            f"Excel 表头必须包含: {REQUIRED_FIELDS}; 实际表头为: {headers}"
        )

    testcases = []
    for row in rows[1:]:
        if all(_to_display_text(cell) == "" for cell in row):
            continue
        record = dict(zip(headers, (_to_display_text(cell) for cell in row)))
        testcases.append(record)

    logger.info("成功加载 %d 条测试用例", len(testcases))
    return testcases


TEST_DATA = load_testcases(TESTCASE_FILE)


@pytest.fixture
def browser():
    """启动 Edge 浏览器并打开登录页面。"""
    options = Options()
    options.add_argument("--start-maximized")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-gpu")

    service = Service(EdgeChromiumDriverManager().install())

    driver = webdriver.Edge(service=service, options=options)
    driver.implicitly_wait(10)

    try:
        driver.get(BASE_URL)
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.ID, "loginBtn"))
        )
        yield driver
    finally:
        driver.quit()


def perform_login(driver, username, password):
    """填充账号密码并点击登录按钮。"""
    wait = WebDriverWait(driver, 15)

    username_input = wait.until(
        EC.presence_of_element_located((By.ID, "username"))
    )
    password_input = wait.until(
        EC.presence_of_element_located((By.ID, "password"))
    )
    login_button = wait.until(
        EC.element_to_be_clickable((By.ID, "loginBtn"))
    )

    username_input.clear()
    username_input.send_keys(str(username))

    password_input.clear()
    password_input.send_keys(str(password))

    original_url = driver.current_url
    logger.info("点击登录按钮，当前地址: %s", original_url)

    login_button.click()
    return original_url


def _is_expected_success(case_type, expected_result):
    """根据用例类型/预期结果判断该用例是否期望登录成功。"""
    case_type = case_type or ""
    expected_result = expected_result or ""

    positive_keywords = ["正向", "正常", "成功", "正确", "有效", "正例", "success", "valid"]
    negative_keywords = ["反向", "失败", "异常", "错误", "无效", "反例", "fail", "invalid", "error"]

    case_type_l = case_type.lower()
    positive_hit = any(k in case_type_l for k in positive_keywords)
    negative_hit = any(k in case_type_l for k in negative_keywords)

    if positive_hit and not negative_hit:
        return True
    if negative_hit and not positive_hit:
        return False

    expected_l = expected_result.lower()
    if any(k in expected_l for k in ["成功", "跳转", "有效", "success", "valid"]):
        return True
    if any(k in expected_l for k in ["失败", "错误", "无效", "异常", "fail", "invalid", "error"]):
        return False

    # 默认按正向成功执行；可通过用例类型明确控制。
    return True


def _get_page_content(driver):
    """获取当前页面源代码。"""
    try:
        return driver.page_source or ""
    except WebDriverException as exc:
        logger.warning("获取页面内容失败: %s", exc)
        return ""


def is_redirect_away_from_login(url, original_url):
    """判断页面是否已经离开登录地址。"""
    if url == original_url:
        return False
    return "login" not in url.lower()


def wait_for_login_response(driver, expected_success, expected_result, original_url):
    """提交登录后等待页面响应，并返回验证信息。"""
    timeout = 12
    start_time = time.time()
    response_text = ""
    final_url = ""
    redirected = False
    expected_found = False

    while time.time() - start_time < timeout:
        try:
            final_url = driver.current_url
            response_text = _get_page_content(driver)
        except WebDriverException as exc:
            logger.debug("获取当前页面状态异常: %s", exc)

        redirected = is_redirect_away_from_login(final_url, original_url)
        expected_found = bool(expected_result) and expected_result in response_text

        if expected_success and (redirected or expected_found):
            break

        if not expected_success:
            if redirected:
                break
            if expected_found:
                break
            if not expected_result and time.time() - start_time >= 1.5:
                break

        if not expected_success and time.time() - start_time >= 2.5:
            break

        time.sleep(0.3)

    # 再留一点时间给服务端渲染/异步刷新
    time.sleep(0.6)

    try:
        final_url = driver.current_url
        response_text = _get_page_content(driver)
    except WebDriverException:
        pass

    redirected = is_redirect_away_from_login(final_url, original_url)
    expected_found = bool(expected_result) and expected_result in response_text

    return final_url, response_text, redirected, expected_found


@pytest.mark.parametrize(
    "case",
    TEST_DATA,
    ids=[str(case.get("用例ID", "")) for case in TEST_DATA],
)
def test_login(case):
    """逐条执行登录测试用例。"""
    case_id = case.get("用例ID", "")
    username = case.get("账号", "")
    password = case.get("密码", "")
    case_type = case.get("用例类型", "")
    expected_result = case.get("预期结果", "")

    expected_success = _is_expected_success(case_type, expected_result)

    logger.info(
        "开始执行用例: %s | 账号: %s | 密码: %s | 场景: %s | 类型: %s | 预期: %s",
        case_id,
        username,
        "******" if password else "(空)",
        case.get("场景描述", ""),
        case_type,
        expected_result,
    )

    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--start-maximized")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-gpu")

    service = Service(EdgeChromiumDriverManager().install())
    driver = webdriver.Edge(service=service, options=options)
    driver.implicitly_wait(10)

    try:
        driver.get(BASE_URL)
        WebDriverWait(driver, 15).until(
            EC.presence_of_element_located((By.ID, "loginBtn"))
        )

        original_url = perform_login(driver, username, password)
        final_url, response_text, redirected, expected_found = wait_for_login_response(
            driver,
            expected_success=expected_success,
            expected_result=expected_result,
            original_url=original_url,
        )

        if expected_success:
            assert (
                redirected or expected_found
            ), (
                f"用例 [{case_id}] 预期登录成功，但页面未跳转且没有找到预期结果文本。"
                f"\n当前URL: {final_url}"
                f"\n预期文本: {expected_result}"
            )
            logger.info("用例 [%s] 登录成功场景验证通过", case_id)
        else:
            assert not redirected, (
                f"用例 [{case_id}] 预期登录失败，但页面发生了跳转。"
                f"\n当前URL: {final_url}"
            )
            if expected_result and not expected_found:
                logger.warning(
                    "用例 [%s] 仍位于登录页面，但未在页面中找到预期错误文本: %s",
                    case_id,
                    expected_result,
                )
            logger.info("用例 [%s] 登录失败场景验证通过", case_id)

    except (TimeoutException, NoSuchElementException, WebDriverException) as exc:
        logger.error("用例 [%s] 执行异常: %s", case_id, exc)
        pytest.fail(f"用例 [{case_id}] 执行异常: {exc}", pytrace=False)

    except AssertionError as exc:
        logger.error("用例 [%s] 断言失败: %s", case_id, exc)
        raise

    finally:
        driver.quit()