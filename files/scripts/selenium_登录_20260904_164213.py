"""登录功能自动化测试脚本"""
import logging
import re
from pathlib import Path
from typing import List

import openpyxl
import pytest
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.edge.service import Service
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from webdriver_manager.microsoft import EdgeChromiumDriverManager

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

LOGIN_URL = "http://127.0.0.1:8080/login"
USERNAME_INPUT_ID = "username"
PASSWORD_INPUT_ID = "password"
LOGIN_BUTTON_ID = "loginBtn"
SCREENSHOT_DIR = Path(__file__).parent / "screenshots"

EXCEL_FILE = Path(
    r"C:\Users\hp\AppData\Local\Temp\gradio\52b6c5cd2cbd25c28c9abcf34c0560cfd492b1966226f58d1326bb28bd54c33d\login_testcases.xlsx"
)

ERROR_HINT_PREFIXES = (
    "预期结果：",
    "预期结果:",
    "提示：",
    "提示:",
    "错误：",
    "错误:",
    "系统提示：",
    "系统提示:",
    "警告：",
    "警告:",
    "登录失败：",
    "登录失败:",
    "登录失败，",
    "失败：",
    "失败:",
)


def normalize_value(value) -> str:
    """将Excel单元格值转换为去掉首尾空白的字符串"""
    if value is None:
        return ""
    return str(value).strip()


def load_test_cases(excel_path: Path) -> List[dict]:
    """从Excel文件读取全部登录测试用例"""
    workbook = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = workbook.active

    headers = [normalize_value(cell) for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    test_cases = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        raw_data = dict(zip(headers, row))
        test_cases.append(
            {
                "id": normalize_value(raw_data.get("用例ID")),
                "account": normalize_value(raw_data.get("账号")),
                "password": normalize_value(raw_data.get("密码")),
                "scenario": normalize_value(raw_data.get("场景描述")),
                "case_type": normalize_value(raw_data.get("用例类型")),
                "expected_result": normalize_value(raw_data.get("预期结果")),
            }
        )

    workbook.close()
    return test_cases


TEST_CASES = load_test_cases(EXCEL_FILE)


@pytest.fixture
def driver():
    """创建Edge浏览器实例，测试结束后自动关闭"""
    options = webdriver.EdgeOptions()
    options.add_argument("--headless")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")

    service = Service(EdgeChromiumDriverManager().install())
    browser = webdriver.Edge(service=service, options=options)
    logger.info("Edge浏览器已启动")

    yield browser

    browser.quit()
    logger.info("Edge浏览器已关闭")


def extract_error_text(expected_result: str) -> str:
    """从预期结果文本中提取页面可能展示的错误片段"""
    text = expected_result.strip()

    for prefix in ERROR_HINT_PREFIXES:
        if text.startswith(prefix):
            text = text[len(prefix):]
            break

    text = re.split(r"[。，；;！!]", text)[0].strip()
    return text


def fill_and_submit_login_form(browser, account: str, password: str):
    """填充用户名密码并点击登录按钮"""
    wait = WebDriverWait(browser, 10)

    username_input = wait.until(
        EC.visibility_of_element_located((By.ID, USERNAME_INPUT_ID))
    )
    password_input = wait.until(
        EC.visibility_of_element_located((By.ID, PASSWORD_INPUT_ID))
    )

    username_input.clear()
    if account:
        username_input.send_keys(account)

    password_input.clear()
    if password:
        password_input.send_keys(password)

    login_button = wait.until(
        EC.element_to_be_clickable((By.ID, LOGIN_BUTTON_ID))
    )
    login_button.click()
    logger.info("已填写账号和密码并点击登录按钮")


def is_still_on_login_page(browser) -> bool:
    """判断当前页面是否仍然停留在登录页"""
    path = browser.current_url.lower().split("?")[0].rstrip("/")
    return path.endswith("/login")


def verify_login_result(browser, case_id: str, expected_result: str):
    """根据预期结果验证登录是否成功或失败"""
    wait = WebDriverWait(browser, 10)

    if "成功" in expected_result:
        wait.until(lambda d: not is_still_on_login_page(d))
        logger.info("用例 %s 登录成功，当前页面URL: %s", case_id, browser.current_url)
        return

    error_text = extract_error_text(expected_result)
    if not error_text:
        logger.warning("用例 %s 未能从预期结果中提取错误提示文本，跳过页面文本校验", case_id)
        return

    wait.until(
        lambda d: error_text in d.find_element(By.TAG_NAME, "body").text
    )
    logger.info("用例 %s 验证失败场景成功，页面包含提示: %s", case_id, error_text)


@pytest.mark.parametrize(
    "case_id, account, password, scenario, case_type, expected_result",
    [
        (
            case["id"],
            case["account"],
            case["password"],
            case["scenario"],
            case["case_type"],
            case["expected_result"],
        )
        for case in TEST_CASES
    ],
    ids=[case["id"] for case in TEST_CASES],
)
def test_login(
    case_id: str,
    account: str,
    password: str,
    scenario: str,
    case_type: str,
    expected_result: str,
    driver,
):
    """执行单条登录测试用例"""
    logger.info("开始执行用例ID=%s, 场景=%s, 类型=%s, 预期=%s",
                case_id, scenario, case_type, expected_result)

    try:
        driver.get(LOGIN_URL)
        fill_and_submit_login_form(driver, account, password)
        verify_login_result(driver, case_id, expected_result)

    except Exception as exc:
        SCREENSHOT_DIR.mkdir(exist_ok=True)
        screenshot_path = SCREENSHOT_DIR / f"{case_id}.png"
        driver.save_screenshot(str(screenshot_path))
        logger.exception("用例 %s 执行失败，截图已保存到 %s", case_id, screenshot_path)
        raise AssertionError(f"用例 {case_id} 执行失败: {exc}") from exc


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v", "--tb=short"]))